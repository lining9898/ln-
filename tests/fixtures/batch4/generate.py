"""Regenerate small synthetic, non-private BATCH 4 fixtures (developer only)."""
from pathlib import Path
import fitz
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt

root = Path(__file__).parent
pdf = fitz.open()
page = pdf.new_page()
page.insert_text((72, 80), 'English first page, physical page 1.', fontsize=16)
page.insert_text((72, 120), '中文第一章：资料来源', fontsize=16, fontname='china-s')
pdf.new_page()  # Physical page 2 intentionally blank.
page = pdf.new_page()
page.insert_text((72, 80), 'English last page, physical page 3.', fontsize=16)
page.insert_text((72, 120), '中文第三页：正文', fontsize=16, fontname='china-s')
pdf.save(root / 'text_pages.pdf')
pdf.close()
image = Image.new('RGB', (1600, 520), 'white')
draw = ImageDraw.Draw(image)
font = ImageFont.truetype('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc', 82)
draw.text((80, 80), '扫描资料 中文测试', font=font, fill='black')
draw.text((80, 230), 'Scanned English 2026', font=font, fill='black')
image.save(root / 'scan.jpg', quality=82)
pdf = fitz.open()
pdf.new_page().insert_text((72, 80), 'Text page before scanned page.', fontsize=16)
pdf.new_page().insert_image(fitz.Rect(25, 80, 570, 257), filename=str(root / 'scan.jpg'))
pdf.save(root / 'mixed_scan.pdf')
pdf.close()
(root / 'scan.jpg').unlink()  # image is embedded in the real PDF
pdf = fitz.open(root / 'mixed_scan.pdf')
scan_only = fitz.open()
scan_only.insert_pdf(pdf, from_page=1, to_page=1)
scan_only.save(root / 'scan_only.pdf')
scan_only.close()
pdf.close()
Document().save(root / 'empty.docx')
word = Document()
word.add_heading('第一章 公司资料', level=1)
word.add_paragraph('中文正文 English paragraph one.')
word.add_heading('第二节 技术说明', level=2)
word.add_paragraph('第二节正文，带空格的项目名称 A.')
word.add_heading('小节 三级标题', level=3)
word.add_paragraph('三级标题后的普通段落。')
table = word.add_table(rows=2, cols=2)
table.cell(0, 0).text = '项目'
table.cell(0, 1).text = '说明'
table.cell(1, 0).text = '测试 A'
table.cell(1, 1).text = '表格中文内容'
word.save(root / 'headings_table.docx')
(root / 'lines 中文.txt').write_text('中文第一行 English\n第二行 有空格\n\n第四行 new paragraph\n', encoding='utf-8')
(root / 'headings 中文.md').write_text(
    '# 第一章 总览\n开头中文内容\n\n## 第二节 细节\n段落 English text\n\n'
    '```text\n# code heading should stay text\n代码内容\n```\n\n'
    '### 第三层 标题\n最后内容\n', encoding='utf-8')
print('Generated:', *[(p.name, p.stat().st_size) for p in root.iterdir() if p.is_file()])
